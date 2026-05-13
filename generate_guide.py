"""
SolidWORKS アセンブリ コピー＆リネーム マクロ
操作説明資料（Excel）生成スクリプト

実行方法:
    pip install openpyxl
    python generate_guide.py
"""

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

# ============================================================
# カラー定義
# ============================================================
COLOR_HEADER_BG    = "1F4E79"   # 濃い青（ヘッダー背景）
COLOR_HEADER_FG    = "FFFFFF"   # 白（ヘッダー文字）
COLOR_STEP_BG      = "BDD7EE"   # 薄い青（手順番号背景）
COLOR_CAUTION_BG   = "FFF2CC"   # 黄（注意事項背景）
COLOR_CAUTION_FG   = "7F6000"   # 濃い黄（注意事項文字）
COLOR_SECTION_BG   = "D6E4F0"   # ライトブルー（セクション）
COLOR_GOOD_BG      = "E2EFDA"   # 薄い緑（OKマーク）
COLOR_TITLE_BG     = "2E75B6"   # 中間青（タイトル）

# ============================================================
# スタイル生成ヘルパー
# ============================================================
def hfont(bold=False, size=11, color="000000", name="游ゴシック"):
    return Font(bold=bold, size=size, color=color, name=name)

def hfill(color):
    return PatternFill("solid", fgColor=color)

def hborder(style="thin"):
    s = Side(style=style)
    return Border(left=s, right=s, top=s, bottom=s)

def halign(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def apply_cell(ws, row, col, value, bold=False, size=11, fg="000000",
               bg=None, h="left", v="center", wrap=False, border=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = hfont(bold=bold, size=size, color=fg)
    c.alignment = halign(h=h, v=v, wrap=wrap)
    if bg:
        c.fill = hfill(bg)
    if border:
        c.border = hborder()
    return c

def merge_and_apply(ws, r1, c1, r2, c2, value, bold=False, size=11,
                    fg="000000", bg=None, h="left", v="center", wrap=False):
    ws.merge_cells(start_row=r1, start_column=c1,
                   end_row=r2, end_column=c2)
    c = ws.cell(row=r1, column=c1, value=value)
    c.font      = hfont(bold=bold, size=size, color=fg)
    c.alignment = halign(h=h, v=v, wrap=wrap)
    if bg:
        c.fill = hfill(bg)
    return c

def set_column_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

# ============================================================
# シート1：概要
# ============================================================
def build_overview(wb):
    ws = wb.active
    ws.title = "概要"
    ws.sheet_view.showGridLines = False

    set_column_widths(ws, [3, 20, 55, 3])

    # タイトル
    merge_and_apply(ws, 1, 2, 2, 3,
                    "SolidWORKS アセンブリ コピー＆リネーム マクロ\n操作説明資料",
                    bold=True, size=16, fg=COLOR_HEADER_FG, bg=COLOR_TITLE_BG,
                    h="center", v="center", wrap=True)
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 20

    r = 4
    # マクロの目的
    apply_cell(ws, r, 2, "■ マクロの目的", bold=True, size=13,
               bg=COLOR_SECTION_BG)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 20
    r += 1

    overview_text = (
        "SolidWORKS の親アセンブリ（.sldasm）と、その子コンポーネントのファイルを\n"
        "まとめて新しい名前でコピーするマクロです。\n\n"
        "コピー後の親アセンブリを開くと、コピーされた子ファイルと正しくリンクされた\n"
        "状態になっています。元のファイルはそのまま残ります。"
    )
    apply_cell(ws, r, 2, overview_text, size=11, wrap=True, v="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 80
    r += 2

    # 動作の仕組み
    apply_cell(ws, r, 2, "■ 動作の仕組み", bold=True, size=13,
               bg=COLOR_SECTION_BG)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 20
    r += 1

    steps = [
        ("STEP 1", "SolidWORKS で親アセンブリを開く"),
        ("STEP 2", "マクロを起動する（Tools → Macro → Run Macro）"),
        ("STEP 3", "ダイアログで新しいアセンブリ名と保存先フォルダを指定して「実行」をクリック"),
        ("STEP 4", "自動的に親ファイルと子ファイルのコピーが作成される"),
        ("STEP 5", "「ファイル名を変更しコピーが完了しました。」と表示されたら完了"),
    ]
    for step, desc in steps:
        apply_cell(ws, r, 2, step, bold=True, fg="FFFFFF", bg=COLOR_HEADER_BG,
                   h="center", border=True)
        apply_cell(ws, r, 3, desc, wrap=True, border=True)
        ws.row_dimensions[r].height = 22
        r += 1
    r += 1

    # 子ファイルの検出ルール
    apply_cell(ws, r, 2, "■ 子ファイルの検出ルール", bold=True, size=13,
               bg=COLOR_SECTION_BG)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 20
    r += 1

    rule_text = (
        "アセンブリ内のコンポーネントのうち、ファイル名が「親ファイル名」で始まるものが\n"
        "コピー対象となります。それ以外のコンポーネントはコピーされません。\n\n"
        "【例】\n"
        "  親ファイル名：WIDGET-A.sldasm\n"
        "  → WIDGET-A-01.sldprt  ✓ コピー対象\n"
        "  → WIDGET-A-FRAME.sldprt  ✓ コピー対象\n"
        "  → OTHER-PART.sldprt  ✗ コピー対象外（名前が違う）"
    )
    apply_cell(ws, r, 2, rule_text, wrap=True, v="top",
               bg=COLOR_CAUTION_BG)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 120
    r += 2

    # 注意事項
    apply_cell(ws, r, 2, "⚠ 実行前の注意事項", bold=True, size=13,
               fg=COLOR_CAUTION_FG, bg=COLOR_CAUTION_BG)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 20
    r += 1

    cautions = [
        "□ 元ファイルのバックアップを取得してください",
        "□ 親アセンブリが SolidWORKS で開いている状態でマクロを実行してください",
        "□ 子ファイルが他のプロセスによってロックされていないことを確認してください",
        "□ 保存先フォルダへの書き込み権限があることを確認してください",
    ]
    for c in cautions:
        apply_cell(ws, r, 2, c, wrap=True, fg=COLOR_CAUTION_FG,
                   bg=COLOR_CAUTION_BG, border=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.row_dimensions[r].height = 18
        r += 1


# ============================================================
# シート2：インストール手順
# ============================================================
def build_install(wb):
    ws = wb.create_sheet("インストール手順")
    ws.sheet_view.showGridLines = False
    set_column_widths(ws, [3, 6, 20, 40, 3])

    merge_and_apply(ws, 1, 2, 1, 4,
                    "マクロのインストール手順",
                    bold=True, size=14, fg=COLOR_HEADER_FG, bg=COLOR_HEADER_BG,
                    h="center", v="center")
    ws.row_dimensions[1].height = 28

    steps = [
        ("1", "マクロファイルを準備する",
         "以下の2つのファイルをパソコン上のフォルダに保存します：\n"
         "  • RenameAndCopy.bas\n"
         "  • frmRenameDialog.frm"),
        ("2", "SolidWORKS のマクロエディタを開く",
         "メニューから：\nTools（ツール）→ Macro（マクロ）→ Edit Macro（マクロを編集）"),
        ("3", "新しいマクロプロジェクトを作成する",
         "「Edit Macro」ダイアログで「新規（New）」をクリックし、\n"
         "保存名を「RenameAndCopy.swp」として保存します。"),
        ("4", "モジュールをインポートする",
         "VBAエディタのメニューから：\nFile → Import File\n"
         "① RenameAndCopy.bas を選択してインポート\n"
         "② frmRenameDialog.frm を選択してインポート"),
        ("5", "UserFormのコントロールを配置する",
         "インポートした frmRenameDialog をダブルクリックして開き、\n"
         "ツールボックスから以下のコントロールを配置してください：\n"
         "  • lblCurrentLabel（Label）：「現在のアセンブリ名：」\n"
         "  • lblCurrentName（Label）：アセンブリ名表示\n"
         "  • lblNewName（Label）：「新しいアセンブリ名：」\n"
         "  • txtNewName（TextBox）：新しい名前入力欄\n"
         "  • lblNewNameNote（Label）：「（拡張子 .sldasm は不要です）」\n"
         "  • lblFolder（Label）：「保存先フォルダ：」\n"
         "  • txtFolder（TextBox）：フォルダパス表示・入力\n"
         "  • btnBrowse（CommandButton）：「参照...」\n"
         "  • lblPreview（Label）：対象子ファイル数表示\n"
         "  • btnExecute（CommandButton）：「実行」\n"
         "  • btnCancel（CommandButton）：「キャンセル」"),
        ("6", "マクロを保存する",
         "VBAエディタで Ctrl+S を押してマクロプロジェクトを保存します。"),
    ]

    r = 3
    for num, title, detail in steps:
        apply_cell(ws, r, 2, num, bold=True, fg="FFFFFF", bg=COLOR_HEADER_BG,
                   h="center", v="top", border=True)
        apply_cell(ws, r, 3, title, bold=True, bg=COLOR_STEP_BG, border=True)
        apply_cell(ws, r, 4, detail, wrap=True, v="top", border=True)
        ws.row_dimensions[r].height = max(18, detail.count("\n") * 15 + 18)
        r += 1

    # 補足
    r += 1
    merge_and_apply(ws, r, 2, r, 4,
                    "💡 補足：SolidWORKS 2020 以降のバージョンで動作確認済みです。",
                    fg=COLOR_CAUTION_FG, bg=COLOR_CAUTION_BG, wrap=True)
    ws.row_dimensions[r].height = 18


# ============================================================
# シート3：操作手順
# ============================================================
def build_operation(wb):
    ws = wb.create_sheet("操作手順")
    ws.sheet_view.showGridLines = False
    set_column_widths(ws, [3, 6, 22, 40, 3])

    merge_and_apply(ws, 1, 2, 1, 4,
                    "マクロの操作手順",
                    bold=True, size=14, fg=COLOR_HEADER_FG, bg=COLOR_HEADER_BG,
                    h="center", v="center")
    ws.row_dimensions[1].height = 28

    r = 3
    sections = [
        ("準備", [
            ("1", "親アセンブリのバックアップを取得する",
             "必ず元ファイルをコピーしてバックアップを取ってください。"),
            ("2", "SolidWORKS で親アセンブリを開く",
             "リネーム対象の親アセンブリ（.sldasm）を SolidWORKS で開きます。\n"
             "パーツファイル（.sldprt）だけを開いた状態では動作しません。"),
        ]),
        ("マクロ実行", [
            ("3", "マクロを起動する",
             "メニューから：\nTools（ツール）→ Macro（マクロ）→ Run Macro（マクロを実行）\n"
             "「RenameAndCopy.swp」を選択して「開く」→ StartRenameAndCopy を選択して「実行」"),
            ("4", "ダイアログが表示される",
             "以下の情報が表示されるダイアログが開きます：\n"
             "  • 現在のアセンブリ名（確認用）\n"
             "  • 対象となる子ファイル数\n"
             "  入力欄に新しいアセンブリ名と保存先フォルダを指定します。"),
            ("5", "新しいアセンブリ名を入力する",
             "「新しいアセンブリ名」欄に新しい名前を入力します。\n"
             "例：WIDGET-A → GADGET-B\n"
             "拡張子（.sldasm）は入力不要です。"),
            ("6", "保存先フォルダを指定する",
             "デフォルトは元ファイルと同じフォルダです。\n"
             "「参照...」ボタンでフォルダを選択するか、直接パスを入力できます。\n"
             "存在しないフォルダを指定した場合は作成するか確認されます。"),
            ("7", "「実行」ボタンをクリックする",
             "確認ダイアログが表示されます。内容を確認して「はい」をクリックすると処理が始まります。"),
        ]),
        ("完了確認", [
            ("8", "完了メッセージを確認する",
             "「ファイル名を変更しコピーが完了しました。」というメッセージが表示されたら完了です。\n"
             "コピーされたファイルの数が表示されます。"),
            ("9", "新しい親アセンブリを開いて確認する",
             "保存先フォルダに作成された新しい親アセンブリ（.sldasm）を開きます。\n"
             "全ての子コンポーネントが正常に表示されることを確認してください。"),
        ]),
    ]

    for section_name, step_list in sections:
        # セクションヘッダー
        merge_and_apply(ws, r, 2, r, 4,
                        f"【 {section_name} 】",
                        bold=True, size=12, fg=COLOR_HEADER_FG, bg=COLOR_TITLE_BG)
        ws.row_dimensions[r].height = 20
        r += 1

        for num, title, detail in step_list:
            apply_cell(ws, r, 2, num, bold=True, fg="FFFFFF", bg=COLOR_HEADER_BG,
                       h="center", v="top", border=True)
            apply_cell(ws, r, 3, title, bold=True, bg=COLOR_STEP_BG, border=True, v="top")
            apply_cell(ws, r, 4, detail, wrap=True, v="top", border=True)
            line_count = detail.count("\n") + 1
            ws.row_dimensions[r].height = max(20, line_count * 16)
            r += 1
        r += 1

    # コピー前後の対応表
    merge_and_apply(ws, r, 2, r, 4,
                    "【 コピー前後のファイル名の例 】",
                    bold=True, size=12, fg=COLOR_HEADER_FG, bg=COLOR_TITLE_BG)
    ws.row_dimensions[r].height = 20
    r += 1

    # ヘッダー行
    apply_cell(ws, r, 2, "種別",     bold=True, fg="FFFFFF", bg=COLOR_HEADER_BG,
               h="center", border=True)
    apply_cell(ws, r, 3, "コピー前（元ファイル）", bold=True, fg="FFFFFF",
               bg=COLOR_HEADER_BG, h="center", border=True)
    apply_cell(ws, r, 4, "コピー後（新ファイル）", bold=True, fg="FFFFFF",
               bg=COLOR_HEADER_BG, h="center", border=True)
    ws.row_dimensions[r].height = 18
    r += 1

    examples = [
        ("親アセンブリ", "WIDGET-A.sldasm",        "GADGET-B.sldasm"),
        ("子パーツ①",   "WIDGET-A-01.sldprt",     "GADGET-B-01.sldprt"),
        ("子パーツ②",   "WIDGET-A-02.sldprt",     "GADGET-B-02.sldprt"),
        ("子パーツ③",   "WIDGET-A-FRAME.sldprt",  "GADGET-B-FRAME.sldprt"),
        ("対象外",       "OTHER-PART.sldprt",       "コピーされない"),
    ]
    for kind, before, after in examples:
        apply_cell(ws, r, 2, kind,   h="center", border=True)
        apply_cell(ws, r, 3, before, border=True)
        bg = COLOR_GOOD_BG if after != "コピーされない" else COLOR_CAUTION_BG
        apply_cell(ws, r, 4, after, bg=bg, border=True)
        ws.row_dimensions[r].height = 18
        r += 1


# ============================================================
# シート4：よくある質問
# ============================================================
def build_faq(wb):
    ws = wb.create_sheet("よくある質問")
    ws.sheet_view.showGridLines = False
    set_column_widths(ws, [3, 8, 55, 3])

    merge_and_apply(ws, 1, 2, 1, 3,
                    "よくある質問（トラブルシューティング）",
                    bold=True, size=14, fg=COLOR_HEADER_FG, bg=COLOR_HEADER_BG,
                    h="center", v="center")
    ws.row_dimensions[1].height = 28

    faqs = [
        ("Q", "マクロを実行すると「アセンブリファイルを開いてください」と表示される",
         "A", "パーツファイル（.sldprt）が開かれています。\n"
              "親アセンブリ（.sldasm）を SolidWORKS で開いてから再度マクロを実行してください。"),
        ("Q", "子ファイルがコピーされない",
         "A", "子ファイルのファイル名が親ファイル名で始まっていない可能性があります。\n"
              "命名規則を確認してください。\n"
              "例：親が「WIDGET-A.sldasm」の場合、子は「WIDGET-A-」で始まる必要があります。"),
        ("Q", "コピー後の親アセンブリを開いたら参照エラーが出る",
         "A", "以下を確認してください：\n"
              "① 子ファイルのコピーが正しく完了しているか\n"
              "② コピー先フォルダに全てのファイルが存在するか\n"
              "問題が解消しない場合は、SolidWORKS の「外部参照の修復」機能を使用して\n"
              "手動でパスを再指定してください。"),
        ("Q", "「既に同名ファイルが存在します」と表示される",
         "A", "指定した保存先に同じ名前のファイルが既に存在します。\n"
              "「はい」を選択すると上書きします。\n"
              "「いいえ」を選択するとそのファイルのみスキップします。\n"
              "「キャンセル」を選択すると以降の全ファイルをスキップします。"),
        ("Q", "サブアセンブリ（入れ子になったアセンブリ）は対応しているか",
         "A", "サブアセンブリ自体のファイル名が親ファイル名で始まる場合はコピーされます。\n"
              "ただし、サブアセンブリの内部に含まれる子部品については、\n"
              "サブアセンブリを開いて再度マクロを実行することをお勧めします。"),
        ("Q", "保存先フォルダが存在しないと言われる",
         "A", "「フォルダを作成しますか？」というダイアログで「はい」を選択すると\n"
              "自動的にフォルダが作成されます。\n"
              "フォルダの作成権限がない場合は管理者に相談してください。"),
        ("Q", "元のファイルが変更・削除されないか心配",
         "A", "このマクロは元のファイルを変更・削除しません。\n"
              "コピー（複製）のみを行います。元ファイルはそのまま残ります。"),
    ]

    r = 3
    for i in range(0, len(faqs), 1):
        q_label, q_text, a_label, a_text = faqs[i]

        # Q
        apply_cell(ws, r, 2, q_label, bold=True, fg="FFFFFF", bg=COLOR_HEADER_BG,
                   h="center", v="top", border=True)
        apply_cell(ws, r, 3, q_text, bold=True, bg=COLOR_STEP_BG, wrap=True,
                   v="top", border=True)
        ws.row_dimensions[r].height = max(20, q_text.count("\n") * 14 + 18)
        r += 1

        # A
        apply_cell(ws, r, 2, a_label, bold=True, fg=COLOR_CAUTION_FG,
                   bg=COLOR_CAUTION_BG, h="center", v="top", border=True)
        apply_cell(ws, r, 3, a_text, wrap=True, v="top", border=True)
        ws.row_dimensions[r].height = max(20, a_text.count("\n") * 14 + 18)
        r += 1
        r += 1  # 間隔

    # フッター
    merge_and_apply(ws, r, 2, r, 3,
                    "その他の問題は VBA エディタの「表示 → イミディエイトウィンドウ」でログを確認してください。",
                    fg=COLOR_CAUTION_FG, bg=COLOR_CAUTION_BG, wrap=True)
    ws.row_dimensions[r].height = 20


# ============================================================
# メイン処理
# ============================================================
def main():
    wb = Workbook()

    build_overview(wb)
    build_install(wb)
    build_operation(wb)
    build_faq(wb)

    output_path = "マクロ使用方法ガイド.xlsx"
    wb.save(output_path)
    print(f"Excel ガイドを生成しました：{output_path}")


if __name__ == "__main__":
    main()
